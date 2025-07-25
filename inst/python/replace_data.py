import sys
import datetime as dt
import numpy as np
import pandas as pd

from pptx import Presentation
from pptx.chart.data import CategoryChartData

from excel_date_format import convert_excel_date_fmt

def py_replace_data(input_file, slide_idx, chart_idx, data, categories, series, output_file, date_fmt = "mm/yy"):
  pres = Presentation(input_file)
  
  if len(pres.slides) < slide_idx:
    raise Exception("Slide " + str(int(slide_idx)) + " does not exist.")

  slide = pres.slides[int(slide_idx - 1)]
  charts = [shape for shape in slide.shapes if shape.has_chart]
  
  if len(charts) < chart_idx:
    raise Exception("Chart " + str(int(chart_idx)) + " in " + " slide " + str(int(slide_idx)) + " does not exist.")
  
  chart = charts[int(chart_idx - 1)].chart
  
  date_fmt = convert_excel_date_fmt(date_fmt)
  date_col = data.map(lambda x: isinstance(x, dt.date)).apply(lambda x: x.all())
  
  if isinstance(categories, str):
    categories = [categories]
    
  if isinstance(series, str):
    series = [series]
    
  data[date_col.index[date_col]] = data[date_col.index[date_col]].map(lambda x: "" if x == dt.date(1900, 1, 1) else x)
  
  if len(categories) > 1:
      data[date_col.index[date_col]] = data[date_col.index[date_col]].map(lambda x: x.strftime(date_fmt) if x != "" else x)
  
  data = data.replace(np.nan, None)
  
  chart_data = CategoryChartData()
  
  if len(categories) > 1:
      for row in data[categories].itertuples(index = False, name = None):
          for i, label in enumerate(row):
              if i == 0:
                  cat = chart_data.add_category(label)
              else:
                  cat = cat.add_sub_category(label)
  else:
      chart_data.categories = data[categories[0]]
  
  for s in series:
      chart_data.add_series(s, data[s])
  
  chart.replace_data(chart_data)
  
  pres.save(output_file)
